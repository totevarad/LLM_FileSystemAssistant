"""
File System Tools for LLM-Powered File System Assistant.

This module provides pure, synchronous, dependency-light Python functions
for interacting with the local file system. It is LLM-agnostic and does not
import or depend on any LLM provider SDKs.

Tools provided:
    - read_file(filepath: str) -> dict
    - list_files(directory: str, extension: Optional[str] = None) -> list
    - write_file(filepath: str, content: str) -> dict
    - search_in_file(filepath: str, keyword: str) -> dict
"""

from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any, Tuple

import docx
import pypdf

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def _read_txt(path: Path) -> Tuple[str, Optional[int]]:
    """Read a .txt file with UTF-8 encoding and latin-1 fallback."""
    raw_bytes = path.read_bytes()
    try:
        content = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = raw_bytes.decode("latin-1", errors="replace")
    return content, None


def _read_pdf(path: Path) -> Tuple[str, Optional[int]]:
    """Read a .pdf file and extract text and page count using pypdf."""
    reader = pypdf.PdfReader(str(path))
    if reader.is_encrypted:
        try:
            # Attempt decrypt with empty password for unauthenticated encrypted PDFs
            reader.decrypt("")
        except Exception as exc:
            raise ValueError(f"Encrypted/password-protected PDF: {exc}") from exc

    pages_text = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            pages_text.append(extracted)
    content = "\n".join(pages_text)
    num_pages = len(reader.pages)
    return content, num_pages


def _read_docx(path: Path) -> Tuple[str, Optional[int]]:
    """Read a .docx file and extract text from paragraphs and tables using python-docx."""
    doc = docx.Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_rows = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_rows.append(row_text)

    all_parts = paragraphs + table_rows
    content = "\n".join(all_parts)
    return content, None


def read_file(filepath: str) -> Dict[str, Any]:
    """Read a document (.txt, .pdf, .docx) and return its extracted text content and metadata.

    Parameters:
        filepath (str): Relative or absolute path to the target file.

    Returns:
        dict: Structured dictionary containing:
            - success (bool): True if reading succeeded, False otherwise.
            - filepath (str): Original path provided.
            - filename (str): Name of the file with extension.
            - extension (str): Lowercase file extension (e.g., '.txt', '.pdf', '.docx').
            - content (Optional[str]): Extracted text content, or None on failure.
            - metadata (Optional[dict]): File statistics including size, page count,
              word count, character count, modified time, and read time. None on failure.
            - error (Optional[str]): Error description if failed, None if successful.
    """
    if not filepath or not isinstance(filepath, str):
        return {
            "success": False,
            "filepath": str(filepath) if filepath is not None else "",
            "filename": "",
            "extension": "",
            "content": None,
            "metadata": None,
            "error": "Invalid filepath: filepath must be a non-empty string.",
        }

    path = Path(filepath)
    filename = path.name
    extension = path.suffix.lower()

    try:
        if not path.exists():
            return {
                "success": False,
                "filepath": filepath,
                "filename": filename,
                "extension": extension,
                "content": None,
                "metadata": None,
                "error": f"File not found: '{filepath}'.",
            }

        if path.is_dir():
            return {
                "success": False,
                "filepath": filepath,
                "filename": filename,
                "extension": extension,
                "content": None,
                "metadata": None,
                "error": f"Path is a directory, not a file: '{filepath}'.",
            }

        if extension not in SUPPORTED_EXTENSIONS:
            return {
                "success": False,
                "filepath": filepath,
                "filename": filename,
                "extension": extension,
                "content": None,
                "metadata": None,
                "error": (
                    f"Unsupported file extension '{extension}'. "
                    f"Supported formats are: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
                ),
            }

        # Extract content and page count based on format
        if extension == ".txt":
            content, num_pages = _read_txt(path)
        elif extension == ".pdf":
            content, num_pages = _read_pdf(path)
        elif extension == ".docx":
            content, num_pages = _read_docx(path)
        else:
            raise ValueError(f"Unhandled extension: {extension}")

        stat = path.stat()
        num_characters = len(content)
        num_words = len(content.split()) if content else 0

        metadata = {
            "size_bytes": stat.st_size,
            "num_pages": num_pages,
            "num_words": num_words,
            "num_characters": num_characters,
            "modified_time": datetime.fromtimestamp(stat.st_mtime).isoformat(),
            "read_time": datetime.now().isoformat(),
        }

        return {
            "success": True,
            "filepath": filepath,
            "filename": filename,
            "extension": extension,
            "content": content,
            "metadata": metadata,
            "error": None,
        }

    except Exception as exc:
        return {
            "success": False,
            "filepath": filepath,
            "filename": filename,
            "extension": extension,
            "content": None,
            "metadata": None,
            "error": f"Error reading file '{filepath}': {str(exc)}",
        }


def list_files(directory: str, extension: Optional[str] = None) -> List[Dict[str, Any]]:
    """List files in the specified directory, optionally filtered by extension."""
    raise NotImplementedError("Phase 3 implementation pending")


def write_file(filepath: str, content: str) -> Dict[str, Any]:
    """Write text content to a file, creating any missing parent directories."""
    raise NotImplementedError("Phase 4 implementation pending")


def search_in_file(filepath: str, keyword: str) -> Dict[str, Any]:
    """Search for occurrences of a keyword within a file and return context snippets."""
    raise NotImplementedError("Phase 5 implementation pending")
